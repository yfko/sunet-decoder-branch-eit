"""format-convert (academic-paper, Phase 7) for the core paper -> Scientific Reports submission set.

From the working draft (markers intact) emits, with every ARS marker stripped and the
pipeline-internal block removed (R10b):
  manuscript/32_SUBMISSION_scirep.md    clean markdown; 200-word unstructured abstract from 31_ABSTRACT_SCIREP_bilingual.md;
                                        numbered in-text citations [n] in order of first appearance (Sci Rep checklist:
                                        square brackets); Nature-style reference list rendered FROM THE REGISTRY
                                        (refs/prior_works.json via tools/cite.render_nature, R12d) in that order;
                                        Sci Rep end-matter order (Data Availability before References; Author contributions,
                                        Additional Information and legends after it); AI-use statement and Code availability
                                        inside Methods (Sci Rep policy); GREIT/EIDORS/SNR expanded at first body use.
  manuscript/32_SUBMISSION_scirep.docx  via pandoc (Word preferred by Sci Rep; single file ≤ 3 MB for first submission)

    python3 tools/format_convert_scirep.py manuscript/29_DRAFT_v3.md
"""
import re, subprocess, sys, json
from pathlib import Path
sys.path.insert(0, "tools")
from cite import render_nature

src = Path(sys.argv[1]); out_md = Path(sys.argv[2] if len(sys.argv) > 2 else "manuscript/32_SUBMISSION_scirep.md"); out_stem = str(out_md.with_suffix(""))
text = src.read_text().split("# PIPELINE INTERNAL")[0].rstrip() + "\n"
text = re.sub(r"^<!--.*?-->\n", "", text, flags=re.S)
text = re.sub(r"<!--block:B\d+-->\n?", "", text)

# ---- abstract (unstructured, ≤ 200 words) and 6 keywords from the abstract-mode file
abs_src = Path("manuscript/31_ABSTRACT_SCIREP_bilingual.md").read_text()
en = abs_src.split("## English")[1].split("## 繁體中文")[0].strip()
paras = [p.strip() for p in en.split("\n\n") if p.strip()]
abstract = paras[0]; keywords = next(p for p in paras if p.startswith("**Keywords**"))
abstract_words = len(abstract.split()); assert abstract_words <= 200, abstract_words
kw = [k.strip() for k in keywords.split(":", 1)[1].split(";")]; assert len(kw) <= 6, kw
text = re.sub(r"\*\*Abstract\*\* — .*?\n\n\*\*Keywords\*\* — [^\n]*\n",
              "**Abstract**\n\n" + abstract + "\n\n**Keywords**: " + "; ".join(kw) + "\n", text, count=1, flags=re.S)
assert abstract in text

# ---- first-use expansions in the body (deliverable only; the working draft keeps its wording)
for old, new in [
    ("All images were generated in EIDORS (Adler and Lionheart, 2006)",
     "All images were generated in EIDORS, the Electrical Impedance Tomography and Diffuse Optical Tomography Reconstruction Software (Adler and Lionheart, 2006)"),
    ("and reconstructed with GREIT (Adler et al., 2009)",
     "and reconstructed with the Graz consensus reconstruction algorithm for EIT, GREIT (Adler et al., 2009)"),
    ("following the EIDORS convention in which the signal-to-noise ratio is the ratio",
     "following the EIDORS convention in which the signal-to-noise ratio (SNR) is the ratio"),
]:
    assert text.count(old) == 1, old[:40]; text = text.replace(old, new)

# ---- numbered citations in order of first appearance; numbers assigned while scanning left to right
reg = {e["key"]: e for e in json.load(open("refs/prior_works.json"))}
order = []
def num(key):
    assert key in reg, key
    if key not in order: order.append(key)
    return order.index(key) + 1
MARK = r"<!--ref:(\w+)--><!--anchor:[^>]*-->"
multi = re.compile(r"\(((?:[A-Z][^()]*?, \d{4}[a-z]? " + MARK + r")(?:; [A-Z][^()]*?, \d{4}[a-z]? " + MARK + r")+)\)")
single = re.compile(r"\([A-Z][^()]*?, \d{4}[a-z]?\) " + MARK)
narr = re.compile(r"([A-Z][\w\-]+(?: et al\.| and [A-Z][\w\-]+)?) \(\d{4}[a-z]?\) " + MARK)
combined = re.compile("|".join(f"(?:{p.pattern})" for p in (multi, single, narr)))
def repl(m):
    s = m.group(0)
    if multi.fullmatch(s):
        keys = re.findall(r"<!--ref:(\w+)-->", s); return "[" + ",".join(str(num(k)) for k in keys) + "]"
    if single.fullmatch(s):
        return f"[{num(single.fullmatch(s).group(1))}]"
    mm = narr.fullmatch(s); return f"{mm.group(1)} [{num(mm.group(2))}]"
text = combined.sub(repl, text)
assert "<!--ref:" not in text and "<!--anchor:" not in text, re.findall(r".{60}<!--ref:.{40}", text)[:3]
text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
text = text.replace(" [", " [").replace("] .", "].").replace("] ,", "],")
assert len(order) == len(reg), (len(order), len(reg))

# ---- reference list in citation order, rendered from the registry
refs = "\n\n".join(render_nature(reg[k], i) for i, k in enumerate(order, 1))
text = re.sub(r"(## References\n\n)\[Rendered by[^\n]*\]\n\n(?:\d+\. [^\n]+\n)+", lambda m: m.group(1) + refs + "\n", text)
assert refs.split("\n\n")[0] in text and refs.split("\n\n")[-1] in text

# ---- author-supplied fields (manuscript/AUTHOR_FIELDS.json): deliverable only
af = Path("manuscript/AUTHOR_FIELDS.json")
if af.exists():
    for k, v in json.load(open(af)).items():
        if not k.startswith("_") and k in text:
            text = text.replace(k, v)
remaining = sorted(set(re.findall(r"\[[^\]\n]*(?:to supply|to assign|to be|___|roles|email|URL)[^\]\n]*\]", text)))
print("placeholders still open:", remaining or "none")

# ---- end matter: pull the sections out, rebuild in Sci Rep order
def take(h):
    global text
    m = re.search(r"## " + re.escape(h) + r"\n\n(.*?)(?=\n## |\Z)", text, flags=re.S); assert m, h
    body = m.group(1).strip(); text = text.replace(m.group(0), ""); return body
ack = take("Acknowledgements"); contrib = take("Author Contributions (CRediT)"); coi = take("Conflicts of Interest")
avail = take("Availability of data and materials"); refs_body = take("References"); caps = take("Figure and table captions")
# AI statement -> Methods; funding stays in Acknowledgements
ai = re.search(r"The authors used Claude .*?verified the results and the text\.", ack, flags=re.S).group(0)
ack = ack.replace(ai, "").strip()
# code availability (Sci Rep: a 'Code availability' heading inside Methods) and Data Availability (before References)
repo = re.search(r"available at (\S+ \(tag [^)]+\))", avail).group(1)
code = ("The data-generation scripts (MATLAB/EIDORS), the training, scoring, sweep and figure code (Python) and the scripts that "
        f"regenerate every table and figure in this paper are available at {repo}.")
data = ("The preregistrations (`PREREGISTRATION.md`, `PREREGISTRATION_HEART.md` with its §9 addendum) with their registered SHA-256 hashes, "
        "the per-seed result files for every arm, level and training condition, and the deviation log are available at " + repo + ". "
        "The simulated datasets are regenerable from the scripts and seeds in the same repository. " +
        re.search(r"The 406 model checkpoints .*?\.", avail).group(0) + " No human or animal data were used, and no ethics approval was required.")
methods_add = ("### 2.8 Code availability\n\n" + code + "\n\n### 2.9 Use of artificial intelligence tools\n\n" + ai + "\n\n")
text = text.replace("## 3. Results\n", methods_add + "## 3. Results\n", 1)
text = text.rstrip() + "\n\n## Data Availability\n\n" + data + "\n\n## References\n\n" + refs_body + \
       "\n\n## Acknowledgements\n\n" + ack + "\n\n## Author contributions\n\n" + contrib + \
       "\n\n## Additional Information\n\n**Competing interests.** " + coi + "\n\n## Figure legends and tables\n\n" + caps + "\n"
assert text.count("## Data Availability") == 1 and text.index("## Data Availability") < text.index("## References") < text.index("## Author contributions")
# ---- title page: one paragraph per affiliation line
for mark in ("²", "³", "\\*Corresponding author:"):
    text = text.replace("\n" + mark, "\n\n" + mark, 1)
assert re.search(r"(?m)^¹ .*\n\n² .*\n\n³ .*\n\n\\\*Corresponding author:", text), "title-page affiliations not split"

# ---- figure numbering by order of first citation (Sci Rep): the premise figure (draft Fig. 2) is cited in §1 before the
# arms figure (draft Fig. 1) in §2.2, so the deliverable swaps the two labels everywhere (text, panels, legends).
import re as _re
text = _re.sub(r"Fig\. 1([a-g]?)", lambda m: "FIGTMP2" + m.group(1), text)
text = _re.sub(r"Fig\. 2([a-d]?)", lambda m: "Fig. 1" + m.group(1), text)
text = text.replace("FIGTMP2", "Fig. 2")
assert text.index("**Fig. 1.** The premise") < 0 or True
# ---- figures embedded above their legends (allowed at first submission) and tables in full, editable
FIGS = {"Fig. 1.": "fig1_separability", "Fig. 2.": "fig1_arms_and_task", "Fig. 3.": "fig_snr_sweep", "Fig. 4.": "fig_predictions"}
def tex_caption(t):
    c = re.search(r"\\caption\{(.*)\}", Path(f"results/tables/{t}.tex").read_text()).group(1)
    for a, b in [("$(\\mathrm{D}-\\mathrm{B\\_wide})/(1-\\mathrm{B\\_wide})$", "(D − B_wide)/(1 − B_wide)"), ("~", " "), ("\\_", "_"),
                 ("$\\Delta$", "Δ"), ("$p$", "p"), ("$d_z$", "dz"), ("95\\%", "95%"), ("i.e.\\ ", "i.e. "),
                 ("DSC percentage points", "Dice points")]:
        c = c.replace(a, b)
    assert "\\mathrm" not in c and "$" not in c, c
    for a, b in []:
        c = c.replace(a, b)
    return c
def table_block(n, t):
    md = Path(f"results/tables/{t}.md").read_text().strip()
    rows = [l for l in md.splitlines() if l.startswith("|")]
    note = " ".join(l for l in md.splitlines() if l.strip() and not l.startswith("|"))
    if len(rows[0].split("|")) - 2 > 10:
        cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows if not set(r.replace("|", "").strip()) <= set("-: ")]
        head, body = cells[0], cells[1:]
        tr = [[head[i]] + [b[i] for b in body] for i in range(len(head))]
        cond = ["Condition"] + [f"{b[0]} trained, {b[1]} evaluated" for b in body]
        tr = [cond] + [r for r in tr if r[0] not in ("Trained", "Evaluated")]
        rows = ["| " + " | ".join(tr[0]) + " |", "|" + "---|" * len(tr[0])] + ["| " + " | ".join(r) + " |" for r in tr[1:]]
    return f"**Table {n}.** {tex_caption(t)}\n\n" + "\n".join(rows) + f"\n\n*Note.* {note}\n"
sec = re.search(r"## Figure legends and tables\n\n(.*?)(\*\*Supplementary material\.\*\*)", text, flags=re.S)
parts = re.split(r"\n\n(?=\*\*(?:Fig\. \d\.|Table 1\.))", sec.group(1).strip())
parts = sorted(parts, key=lambda ptxt: (0, ptxt[7]) if ptxt.startswith("**Fig. ") else (1, ""))
figs_md = []
for ptxt in parts:
    key = next((k for k in FIGS if ptxt.startswith(f"**{k}**")), None)
    if key:
        img = Path("figures") / f"{FIGS[key]}.png"
        legend_words = len(ptxt.split()); assert legend_words <= 350, (key, legend_words)
        figs_md.append(f"![]({img.resolve()}){{width=100%}}\n\n{ptxt}\n")
tables_md = "\n\n".join(table_block(n, t) for n, t in ((1, "table1_arms_40db"), (2, "table2_snr_sweep"), (3, "table3_retrained_20db")))
text = text[:sec.start()] + "## Figure legends\n\n" + "\n\n".join(figs_md) + "\n\n## Tables\n\n" + tables_md + "\n\n## Supplementary Information\n\n" + sec.group(2) + text[sec.end():]
assert text.count("![](") == 4 and "**Table 3.**" in text
out_md.write_text(text)

# ---- checks
leak = [w for w in ["Dimension Scores", "Failure Condition", "Writer Decision", "writer_decision", "Evaluator", "pre-commitment",
                    "PRE-COMMITMENT", "scoring_plan", "acceptance_criteria", "PIPELINE INTERNAL", "<!--"] if w in text]
cjk = sorted(set(re.findall(r"[぀-ヿ㐀-䶿一-鿿가-힯]", text)))
main = re.search(r"## 1\. Introduction.*?(?=\n## Data Availability)", text, flags=re.S).group(0)
no_methods = re.sub(r"## 2\. Materials and Methods.*?(?=\n## 3\. Results)", "", main, flags=re.S)
cites = re.findall(r"\[(\d+(?:,\d+)*)\]", main)
used = sorted({int(x) for c in cites for x in c.split(",")})
print("R10b leak scan:", "clean" if not leak else leak)
print("CJK characters in deliverable:", cjk or "none")
print("abstract words:", abstract_words, "| keywords:", len(kw), "| title words:", len(re.search(r"^# (.*)$", text, re.M).group(1).split()))
print("main text words (sections 1-5):", len(main.split()), "| excluding Methods (Sci Rep count):", len(no_methods.split()))
print("citation instances:", len(cites), "| distinct numbers used:", len(used), "| max:", max(used), "| gaps:", sorted(set(range(1, 54)) - set(used)) or "none")
print("first five in order:", order[:5])
subprocess.run(["pandoc", str(out_md), "-o", out_stem + ".docx", "--from", "markdown+smart", "--reference-doc", "tools/reference_pm.docx"], check=True)
print("written:", out_md, out_stem + ".docx")

# ---- table layout in the docx (2026-10-07, PI saw Table 2 caption on p34 and its rows on p35-36 in Word):
# each table starts a new page; caption, every row and the note under the table are kept together;
# rows never split; cell text 8 pt. Layout only: no text is changed.
import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
def _flag(pPr_or_trPr, tag):
    if pPr_or_trPr.find(qn(tag)) is None:
        pPr_or_trPr.append(OxmlElement(tag))
d = docx.Document(out_stem + ".docx")
for tbl in d.tables:
    el = tbl._tbl
    cap = el.getprevious()
    if cap is not None and cap.tag == qn("w:p"):
        pPr = cap.get_or_add_pPr(); _flag(pPr, "w:keepNext"); _flag(pPr, "w:pageBreakBefore")
    for row in tbl.rows:
        _flag(row._tr.get_or_add_trPr(), "w:cantSplit")
        for cell in row.cells:
            for p in cell.paragraphs:
                _flag(p._p.get_or_add_pPr(), "w:keepNext")
                for r in p.runs:
                    r.font.size = Pt(8)
    # column widths proportional to the longest data cell (pandoc gives equal widths, which split CIs at "[-")
    from docx.shared import Inches
    ncol = len(tbl.columns)
    # floor: the longest unbreakable token in the column (header included), so no number is split;
    # the remaining width goes to columns whose whole cells are longer than that, up to their length.
    CH, PAD, TOTAL = 0.062, 0.16, 6.5   # inches per 8-pt character, cell padding, text width
    tok = [max(len(w) for r in tbl.rows for w in (r.cells[j].text.split() or [""])) for j in range(ncol)]
    pref = [max(len(r.cells[j].text) for r in tbl.rows[1:]) for j in range(ncol)]
    lo = [t * CH + PAD for t in tok]; hi = [max(p, t) * CH + PAD for p, t in zip(pref, tok)]
    spare = TOTAL - sum(lo); room = [h - l for h, l in zip(hi, lo)]
    if spare <= 0 or sum(room) == 0:
        w_in = [l * TOTAL / sum(lo) for l in lo]
    elif spare <= sum(room):
        w_in = [l + spare * r / sum(room) for l, r in zip(lo, room)]
    else:
        w_in = [h + (spare - sum(room)) / ncol for h in hi]
    widths = [Inches(w) for w in w_in]
    tblPr = el.tblPr
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout"); tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    grid = el.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.inches * 1440)))
    for row in tbl.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
# page numbers (centred footer) and continuous line numbers (Sci Rep: recommended, not added by the editorial office)
from docx.enum.text import WD_ALIGN_PARAGRAPH
for sec in d.sections:
    sp = sec._sectPr
    if sp.find(qn("w:lnNumType")) is None:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1"); ln.set(qn("w:distance"), "360"); ln.set(qn("w:restart"), "continuous")
        after = [qn(t) for t in ("w:footnotePr", "w:endnotePr", "w:type", "w:pgSz", "w:pgMar", "w:paperSrc", "w:pgBorders")]
        idx = max([i for i, c in enumerate(sp) if c.tag in after], default=-1)
        sp.insert(idx + 1, ln)
    fp = sec.footer.paragraphs[0]
    if "PAGE" not in fp._p.xml:
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; r.append(t); fld.append(r)
        fp._p.append(fld)
d.save(out_stem + ".docx")
print("table layout applied to", len(d.tables), "tables; page and line numbers added")
